"""La compilación rechaza muestras y verifica hashes de los archivos fuente."""
from datetime import UTC, datetime, timedelta
from decimal import Decimal
import hashlib
import io

import pytest
from pypdf import PdfReader, PdfWriter
from docx import Document
from sqlalchemy import select

from app.core.config import get_settings
from app.core.exceptions import ConflictError
from app.models.billing import ReportOrder, ReportTariff
from app.models.evaluation_bundle import EvaluationBundle
from app.models.report import ReportRun
from app.models.study import Study
from app.models.survey import Survey
from app.services.consolidated_bundle_service import combine_paid_bundle_reports


@pytest.mark.asyncio
async def test_merge_only_paid_integrity_checked_pdf(
    session, seed_project, seed_user, seed_instrument_draft, tmp_path, monkeypatch,
):
    (_, version) = seed_instrument_draft
    reports_dir = tmp_path / "reports_storage"
    reports_dir.mkdir()
    monkeypatch.setattr(get_settings(), "export_storage_dir", str(tmp_path / "exports_storage"))

    survey = Survey(
        project_id=seed_project.id, instrument_version_id=version.id,
        created_by_user_id=seed_user.id, name="Encuesta", survey_type="CUSTOM",
    )
    session.add(survey)
    await session.flush()

    tariff = ReportTariff(
        min_workers=1, price=Decimal("20.00"), updated_by_user_id=seed_user.id,
    )
    session.add(tariff)
    await session.flush()

    studies = []
    for index in (1, 2):
        study = Study(
            project_id=seed_project.id, survey_id=survey.id,
            instrument_version_id=version.id, name=f"Área {index}",
            study_type="CUSTOM", status="CLOSED",
        )
        session.add(study)
        await session.flush()
        order = ReportOrder(
            study_id=study.id, organization_id=seed_project.organization_id,
            tariff_id=tariff.id, requested_by_user_id=seed_user.id,
            workers=15, tariff_min_workers=1, amount=Decimal("20.00"),
            status="CONSUMED", expires_at=datetime.now(UTC) + timedelta(days=30),
        )
        session.add(order)
        await session.flush()
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        stream = io.BytesIO()
        writer.write(stream)
        pdf_bytes = stream.getvalue()
        pdf_file = reports_dir / f"source-{index}.pdf"
        pdf_file.write_bytes(pdf_bytes)
        session.add(ReportRun(
            study_id=study.id, billing_order_id=order.id,
            status="COMPLETED", output_format="PDF",
            storage_path=str(pdf_file), data_hash=hashlib.sha256(pdf_bytes).hexdigest(),
        ))
        studies.append(study)
    bundle = EvaluationBundle(
        project_id=seed_project.id, organization_id=seed_project.organization_id,
        created_by_user_id=seed_user.id, label="Batería aprobada",
        study_ids=[study.id for study in studies], status="PUBLISHED",
    )
    session.add(bundle)
    await session.commit()

    combined = await combine_paid_bundle_reports(session, bundle, seed_user.id)
    with open(combined.storage_path, "rb") as stream:
        assert len(PdfReader(stream).pages) == 3  # 1 portada ejecutiva + 2 originales
    assert len(combined.manifest["reports"]) == 2
    assert combined.manifest["executive_cover_pages"] == 1
    assert combined.manifest["reports"][0]["start_page"] == 2
    assert combined.manifest["reports"][1]["start_page"] == 3

    # Los informes PDF renderizados desde un DOCX oficial también deben
    # poder consolidarse, verificando la integridad del archivo original.
    document = Document()
    document.add_heading("Informe profesional", level=1)
    document.add_paragraph("Resultado agregado metodológicamente validado.")
    docx_path = reports_dir / "official-source-1.docx"
    document.save(docx_path)
    run = (await session.execute(
        select(ReportRun).where(ReportRun.study_id == studies[0].id)
    )).scalar_one()
    run.output_format = "DOCX"
    run.storage_path = str(docx_path)
    run.data_hash = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    await session.commit()

    hybrid = await combine_paid_bundle_reports(session, bundle, seed_user.id)
    with open(hybrid.storage_path, "rb") as stream:
        assert len(PdfReader(stream).pages) == 3  # 1 portada ejecutiva + 2 originales
    assert hybrid.manifest["reports"][0]["source_format"] == "DOCX"

    # Cualquier manipulación posterior invalida el documento fuente.
    docx_path.write_bytes(b"DOCX alterado")
    with pytest.raises(ConflictError):
        await combine_paid_bundle_reports(session, bundle, seed_user.id)
